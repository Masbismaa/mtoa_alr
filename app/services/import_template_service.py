"""File template import link (Excel), tampilannya disamain sama laporan export (kop, warna, font):
- Links: tempat ngisi. Header di bawah kop, Category & Visibility pake dropdown, Port cuma angka, siap print
- Guide: penjelasan tiap kolom + contoh + aturan import
- Categories: jalur kategori yg bisa dipake + field apa yg dipake/wajib (sumber dropdown Category)
Sengaja ga ada baris contoh di sheet Links: kalau lupa dihapus, ikut ke-import. Contohnya di sheet Guide.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.access_entry_service import get_category_rule
from app.services.category_service import build_category_label, list_usable_category
from app.services.export_service import (
    BAND_FILL,
    CELL_BORDER,
    COLOR_GRAY_TEXT,
    COLOR_RED,
    DETAIL_TEXT_ALIGNMENT,
    REPORT_SUBTITLE,
    make_font,
    setup_print,
    write_cell,
    write_letterhead,
    write_section_title,
    write_table_header,
)
from app.services.import_service import IMPORT_COLUMN_LIST
from app.utils.constants import IMPORT_MAX_ROW_COUNT

HEADER_ROW = 5
FIRST_DATA_ROW = HEADER_ROW + 1
FIELD_LABEL_DICT = {"url": "URL", "address": "Address", "port": "Port"}
RULE_LIST = [
    "Fill one link per row on the Links sheet, starting right below the header.",
    f"Up to {IMPORT_MAX_ROW_COUNT} links per file. Empty rows are skipped.",
    "Every row is checked with the same rules as Add Link. If any row has an error, nothing is imported; "
    "fix the rows listed on screen and upload again.",
    "A link whose URL (or Address + Port) already exists, or appears twice in the file, is rejected.",
    "You become the owner of every imported link.",
    "Access notes are encrypted after import, but this file is not. Delete it once the import is done.",
]


def write_sheet_note(sheet, row, last_column, text):
    """Satu baris keterangan abu-abu selebar tabel."""
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    write_cell(sheet, row, 1, text, font=make_font(9, is_italic=True, color=COLOR_GRAY_TEXT),
               alignment=Alignment(vertical="center"))


def build_link_sheet(sheet, category_count):
    """Sheet isian: kop, header (+ catatan di tiap judul), 500 baris bergaris, dropdown & batas angka."""
    last_column = len(IMPORT_COLUMN_LIST)
    last_row = FIRST_DATA_ROW + IMPORT_MAX_ROW_COUNT - 1
    sheet.title = "Links"
    sheet.sheet_properties.tabColor = COLOR_RED
    sheet.sheet_view.showGridLines = False
    write_letterhead(sheet, 1, last_column, "Access Link Import", f"{REPORT_SUBTITLE}  |  Import template")
    write_sheet_note(sheet, 4, last_column,
                     "Fill one link per row below. * = required. Hover a column title for help, see the Guide sheet for examples.")
    write_table_header(sheet, HEADER_ROW, 1, [column.header + (" *" if column.is_required else "") for column in IMPORT_COLUMN_LIST])

    for index, column in enumerate(IMPORT_COLUMN_LIST, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = column.width
        header_cell = sheet.cell(row=HEADER_ROW, column=index)
        header_cell.comment = Comment(f"{column.guide}\nExample: {column.example}", "ALR", width=260, height=90)
        if column.is_required:
            header_cell.font = make_font(10, is_bold=True, color="FFFFFF")
        # kolom teks dibikin format Text biar Excel ga ngubah isian (misal 0812... / tanggal)
        number_format = "0" if column.key == "port" else "@"
        for row in range(FIRST_DATA_ROW, last_row + 1):
            cell = sheet.cell(row=row, column=index)
            cell.border = CELL_BORDER
            cell.number_format = number_format
            cell.alignment = DETAIL_TEXT_ALIGNMENT if column.key != "port" else Alignment(horizontal="center", vertical="top")
            if (row - FIRST_DATA_ROW) % 2:
                cell.fill = BAND_FILL

    column_letter_dict = {column.key: get_column_letter(index) for index, column in enumerate(IMPORT_COLUMN_LIST, start=1)}
    if category_count:
        category_list_validation = DataValidation(
            type="list", formula1=f"=Categories!$A${FIRST_DATA_ROW}:$A${FIRST_DATA_ROW + category_count - 1}",
            allow_blank=True, showErrorMessage=True, errorTitle="Unknown category",
            error="Pick a category from the list (see the Categories sheet).",
        )
        add_validation(sheet, category_list_validation, column_letter_dict["category"], last_row)
    visibility_validation = DataValidation(
        type="list", formula1='"Private,Public"', allow_blank=True, showErrorMessage=True,
        errorTitle="Invalid visibility", error="Choose Private or Public.",
    )
    add_validation(sheet, visibility_validation, column_letter_dict["visibility"], last_row)
    port_validation = DataValidation(
        type="whole", operator="between", formula1="1", formula2="65535", allow_blank=True, showErrorMessage=True,
        errorTitle="Invalid port", error="Port must be a number from 1 to 65535.",
    )
    add_validation(sheet, port_validation, column_letter_dict["port"], last_row)

    sheet.freeze_panes = f"A{FIRST_DATA_ROW}"
    setup_print(sheet, last_column, last_row, "landscape", title_row=HEADER_ROW)


def add_validation(sheet, validation, column_letter, last_row):
    sheet.add_data_validation(validation)
    validation.add(f"{column_letter}{FIRST_DATA_ROW}:{column_letter}{last_row}")


def write_body_row(sheet, row, value_list, is_band):
    for offset, value in enumerate(value_list):
        write_cell(sheet, row, 1 + offset, value, font=make_font(10), border=CELL_BORDER,
                   alignment=DETAIL_TEXT_ALIGNMENT, fill=BAND_FILL if is_band else None)


def build_guide_sheet(sheet):
    """Panduan: tabel kolom (wajib / cara isi / contoh) + aturan import."""
    sheet.title = "Guide"
    sheet.sheet_view.showGridLines = False
    for letter, width in zip("ABCD", [18, 11, 62, 26]):
        sheet.column_dimensions[letter].width = width
    write_letterhead(sheet, 1, 4, "How to Fill the Template", f"{REPORT_SUBTITLE}  |  Import guide")
    write_table_header(sheet, HEADER_ROW, 1, ["Column", "Required", "What to fill", "Example"])
    for offset, column in enumerate(IMPORT_COLUMN_LIST):
        # contoh angka (Port) ditulis sebagai angka, biar Excel ga ngasih tanda "number stored as text"
        example = int(column.example) if column.example.isdigit() else column.example
        write_body_row(sheet, FIRST_DATA_ROW + offset,
                       [column.header, "Yes" if column.is_required else "No", column.guide, example], offset % 2)
        if column.is_required:
            sheet.cell(row=FIRST_DATA_ROW + offset, column=2).font = make_font(10, is_bold=True, color=COLOR_RED)
    rule_row = FIRST_DATA_ROW + len(IMPORT_COLUMN_LIST) + 1
    write_section_title(sheet, rule_row, 1, "Import rules")
    for offset, rule_text in enumerate(RULE_LIST, start=1):
        sheet.merge_cells(start_row=rule_row + offset, start_column=1, end_row=rule_row + offset, end_column=4)
        # sel gabungan ga bisa auto-tinggi di Excel: tinggi dihitung dari panjang teks (± 140 huruf per baris)
        sheet.row_dimensions[rule_row + offset].height = 16 * (len(rule_text) // 140 + 1)
        write_cell(sheet, rule_row + offset, 1, f"{offset}.  {rule_text}", font=make_font(10),
                   alignment=Alignment(wrap_text=True, vertical="top"))
    setup_print(sheet, 4, rule_row + len(RULE_LIST), "portrait", title_row=HEADER_ROW)


def describe_category_rule(category):
    """Field yg dipake & yg wajib buat kategori ini (ngikut kategori utamanya). Return (dipake, wajib)."""
    rule = get_category_rule(category)
    used_list = [FIELD_LABEL_DICT[field] for field in rule["field_list"]]
    if rule["has_credential_field"]:
        used_list += ["Username", "Access Note"]
    required_list = [FIELD_LABEL_DICT[field] for field in rule["required_field_list"]]
    return ", ".join(used_list) or "Title & Description only", ", ".join(required_list) or "-"


def build_category_sheet(sheet, category_list):
    """Daftar kategori aktif (sumber dropdown di sheet Links) + field yg dipake/wajib."""
    sheet.title = "Categories"
    sheet.sheet_view.showGridLines = False
    for letter, width in zip("ABC", [34, 44, 18]):
        sheet.column_dimensions[letter].width = width
    write_letterhead(sheet, 1, 3, "Categories", f"{REPORT_SUBTITLE}  |  Values for the Category column")
    write_table_header(sheet, HEADER_ROW, 1, ["Category", "Fields used", "Also required"])
    for offset, category in enumerate(category_list):
        used_text, required_text = describe_category_rule(category)
        write_body_row(sheet, FIRST_DATA_ROW + offset, [build_category_label(category), used_text, required_text], offset % 2)
    sheet.freeze_panes = f"A{FIRST_DATA_ROW}"
    setup_print(sheet, 3, FIRST_DATA_ROW + max(len(category_list), 1) - 1, "portrait", title_row=HEADER_ROW)


def build_import_template():
    """File template siap download. Return BytesIO."""
    category_list = list_usable_category()
    workbook = Workbook()
    workbook.properties.title = "Access Link Import Template"
    build_link_sheet(workbook.active, len(category_list))
    build_guide_sheet(workbook.create_sheet())
    build_category_sheet(workbook.create_sheet(), category_list)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
