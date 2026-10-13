"""Import link dari Excel: download template, upload, cek SEMUA baris dulu, baru simpen semua (atau ga sama sekali).

Tiap baris dicek pake aturan yg sama persis kayak form Add Link (prepare_entry_data), termasuk duplikat ke database.
Duplikat antar baris di file yg sama dicek di sini (belum ada di database). Link hasil import milik yg upload.
"""
import io
from collections import namedtuple

from openpyxl import load_workbook

from app.extensions import db
from app.models import AccessEntry
from app.security.file_validator import detect_content_type
from app.services.access_entry_service import (
    AUDIT_ENTITY_TYPE,
    apply_entry_data,
    get_category_rule,
    build_entry_audit_dict,
    prepare_entry_data,
)
from app.services.audit_service import log_audit
from app.services.category_service import build_category_label, list_usable_category
from app.utils.constants import (
    ATTACHMENT_MAX_SIZE_BYTES,
    AUDIT_ACTION_CREATE,
    AUDIT_ACTION_IMPORT,
    CATEGORY_PATH_SEPARATOR,
    IMPORT_MAX_ROW_COUNT,
    VISIBILITY_PRIVATE,
    VISIBILITY_PUBLIC,
)
from app.utils.exceptions import ValidationError, build_error

# kolom template. header dicocokin tanpa beda huruf besar/kecil & tanda * (penanda wajib di template).
# width/guide/example cuma buat tampilan template (import_template_service.py)
ImportColumn = namedtuple("ImportColumn", ["header", "key", "is_required", "width", "guide", "example"])
IMPORT_COLUMN_LIST = [
    ImportColumn("Title", "title", True, 30, "Name of the link (max 150 characters).", "HR Portal"),
    ImportColumn("Category", "category", True, 26, "Pick from the dropdown. The Categories sheet shows which fields each one uses.", "Web > SAP"),
    ImportColumn("URL", "url", False, 36, "Starts with http:// or https://. Required for Web.", "https://hr.spindo.com"),
    ImportColumn("Address", "address", False, 20, "IP address or host name. Required for Network.", "10.0.0.1"),
    ImportColumn("Port", "port", False, 9, "Number from 1 to 65535.", "22"),
    ImportColumn("Username", "username", False, 20, "Login username (not used by General).", "hr.admin"),
    ImportColumn("Access Note", "access_note", False, 26, "Password or access hint. Stored encrypted after import (not used by General).", "(your password)"),
    ImportColumn("Description", "description", False, 40, "Short note about the link.", "Payroll & leave portal"),
    ImportColumn("Visibility", "visibility", False, 12, "Private (only you) or Public (all users, read-only). Empty = Private.", "Private"),
]
REQUIRED_HEADER_LIST = [column.key for column in IMPORT_COLUMN_LIST if column.is_required]
# kop template makan beberapa baris di atas header, jadi header dicari di sekian baris pertama
HEADER_SEARCH_ROW_COUNT = 10
VISIBILITY_BY_TEXT_DICT = {"": VISIBILITY_PRIVATE, "private": VISIBILITY_PRIVATE, "public": VISIBILITY_PUBLIC}
# error yg ditampilin dibatesin biar halaman ga kepanjangan; sisanya cukup disebut jumlahnya
MAX_SHOWN_ERROR_COUNT = 50


def normalize_category_path(text):
    """'Web > SAP', 'web › sap ' -> 'web › sap' (huruf kecil, pemisah disamain) buat dicocokin ke jalur kategori."""
    part_list = [part.strip() for part in text.replace("›", ">").split(">")]
    return CATEGORY_PATH_SEPARATOR.join(part for part in part_list if part).lower()


def build_category_dict():
    """Jalur kategori aktif (huruf kecil) -> Category, misal 'web › sap' -> SAP."""
    return {normalize_category_path(build_category_label(category)): category for category in list_usable_category()}


def normalize_header(value):
    """Judul kolom: 'Title *' / ' title' -> 'title'."""
    return cell_text(value).rstrip("*").strip().lower()


def cell_text(value):
    """Isi sel jadi teks: 22.0 -> '22', kosong -> ''."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_import_row_list(file_storage):
    """Baca file upload. Return [(nomor baris Excel, {key: teks})]. File salah format -> ValidationError."""
    filename = file_storage.filename or ""
    content_bytes = file_storage.stream.read(ATTACHMENT_MAX_SIZE_BYTES + 1)
    if not filename.lower().endswith(".xlsx"):
        raise ValidationError([build_error("import_file", "Upload the .xlsx file made from the template")])
    if len(content_bytes) > ATTACHMENT_MAX_SIZE_BYTES:
        raise ValidationError([build_error("import_file", f"The file is larger than {ATTACHMENT_MAX_SIZE_BYTES // 1048576}MB")])
    try:
        # cek isi beneran xlsx (bukan file lain yg diganti namanya), tanpa makro, ga bisa jadi zip bomb
        detect_content_type(filename, content_bytes)
    except ValueError as error:
        raise ValidationError([build_error("import_file", f"File rejected: {error}")]) from None
    key_by_header_dict = {column.header.lower(): column.key for column in IMPORT_COLUMN_LIST}
    try:
        sheet = load_workbook(io.BytesIO(content_bytes), read_only=True, data_only=True).worksheets[0]
        row_iter = enumerate(sheet.iter_rows(values_only=True), start=1)
        column_key_list = first_key_list = None
        for row_number, raw_row in row_iter:
            key_list = [key_by_header_dict.get(normalize_header(value)) for value in raw_row]
            if first_key_list is None and any(cell_text(value) for value in raw_row):
                first_key_list = key_list
            if all(key in key_list for key in REQUIRED_HEADER_LIST):
                column_key_list = key_list
                break
            if row_number >= HEADER_SEARCH_ROW_COUNT:
                break
        if column_key_list is None:
            found_key_list = first_key_list or []
            missing_header_list = [column.header for column in IMPORT_COLUMN_LIST
                                   if column.is_required and column.key not in found_key_list]
            raise ValidationError([build_error("import_file", f"Missing column(s): {', '.join(missing_header_list)}. Use the template.")])

        row_list = []
        for row_number, raw_row in row_iter:
            row_dict = {key: cell_text(value) for key, value in zip(column_key_list, raw_row) if key}
            if not any(row_dict.values()):
                continue
            if len(row_list) >= IMPORT_MAX_ROW_COUNT:
                raise ValidationError([build_error("import_file", f"Up to {IMPORT_MAX_ROW_COUNT} links per import; split the file.")])
            row_list.append((row_number, row_dict))
    except (ValueError, KeyError, OSError, IndexError):
        raise ValidationError([build_error("import_file", "The file can't be read as an Excel template")]) from None
    if not row_list:
        raise ValidationError([build_error("import_file", "The file has no links to import")])
    return row_list


def build_row_data(row_dict, category_dict):
    """Isi satu baris jadi data form link + daftar error yg ketauan sebelum dicek aturan link."""
    error_list = []
    category = category_dict.get(normalize_category_path(row_dict.get("category", "")))
    if category is None:
        error_list.append(f'Unknown or inactive category "{row_dict.get("category", "")}" (see the Categories sheet)')
    else:
        error_list.extend(find_unused_column_error_list(row_dict, category))
    visibility = VISIBILITY_BY_TEXT_DICT.get(row_dict.get("visibility", "").lower())
    if visibility is None:
        error_list.append("Visibility must be Public or Private")
    data_dict = {column.key: row_dict.get(column.key, "") for column in IMPORT_COLUMN_LIST
                 if column.key not in ("category", "visibility")}
    data_dict.update(category_id=category.id if category else None, visibility=visibility)
    return data_dict, error_list


def find_unused_column_error_list(row_dict, category):
    """Kolom keisi tapi ga dipake kategori ini (misal Access Note di General). Aturan link bakal ngebuang isinya
    diem-diem, jadi di import ditolak biar password/URL yg diisi user ga ilang tanpa ketauan."""
    rule = get_category_rule(category)
    used_key_set = set(rule["field_list"]) | ({"username", "access_note"} if rule["has_credential_field"] else set())
    return [
        f"{column.header} is not used by category {build_category_label(category)}"
        for column in IMPORT_COLUMN_LIST
        if column.key in ("url", "address", "port", "username", "access_note")
        and column.key not in used_key_set and row_dict.get(column.key)
    ]


def find_duplicate_in_file(clean_dict, seen_dict, row_number):
    """Cek URL / Address+Port udah dipake baris lain di file ini. Return pesan error atau None."""
    key_dict = {}
    if clean_dict.get("url"):
        key_dict["URL"] = ("url", clean_dict["url"])
    if clean_dict.get("address"):
        key_dict["Address/Port"] = ("address", clean_dict["address"], clean_dict.get("port"))
    for label, key in key_dict.items():
        if key in seen_dict:
            return f"Same {label} as row {seen_dict[key]}"
    for key in key_dict.values():
        seen_dict[key] = row_number
    return None


def import_entries(user, file_storage):
    """Import semua baris. Ada 1 baris salah -> ga ada yg disimpen, ValidationError isinya error per baris.
    Semua bener -> disimpen dalam 1 transaksi + dicatat (per link & 1 catatan import). Return jumlah link."""
    row_list = read_import_row_list(file_storage)
    category_dict = build_category_dict()
    prepared_list = []
    error_text_list = []
    seen_dict = {}
    for row_number, row_dict in row_list:
        data_dict, row_error_list = build_row_data(row_dict, category_dict)
        if not row_error_list:
            try:
                clean_dict, clean_custom_field_list, _ = prepare_entry_data(user, data_dict, [])
            except ValidationError as error:
                row_error_list = [error_dict["message"] for error_dict in error.error_list]
            else:
                duplicate_text = find_duplicate_in_file(clean_dict, seen_dict, row_number)
                if duplicate_text:
                    row_error_list = [duplicate_text]
                else:
                    prepared_list.append((clean_dict, clean_custom_field_list))
        error_text_list.extend(f"Row {row_number}: {message}" for message in row_error_list)

    if error_text_list:
        shown_list = error_text_list[:MAX_SHOWN_ERROR_COUNT]
        if len(error_text_list) > MAX_SHOWN_ERROR_COUNT:
            shown_list.append(f"...and {len(error_text_list) - MAX_SHOWN_ERROR_COUNT} more errors")
        raise ValidationError([build_error("import_file", text) for text in shown_list])

    try:
        for clean_dict, clean_custom_field_list in prepared_list:
            entry = AccessEntry(user_id=user.id)
            apply_entry_data(entry, clean_dict, clean_custom_field_list)
            db.session.add(entry)
            db.session.flush()
            log_audit(AUDIT_ACTION_CREATE, AUDIT_ENTITY_TYPE, entity_id=entry.id, new_data_dict=build_entry_audit_dict(entry), user=user)
        log_audit(AUDIT_ACTION_IMPORT, AUDIT_ENTITY_TYPE, user=user,
                  new_data_dict={"row_count": len(prepared_list), "filename": (file_storage.filename or "")[:255]})
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return len(prepared_list)
