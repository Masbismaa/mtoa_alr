from collections import Counter
from io import BytesIO
from functools import lru_cache
from pathlib import Path
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from app.services.category_service import build_category_label, get_root_category
from app.utils.chart_helper import calculate_percent
from app.utils.constants import (
    EXPORT_FORMULA_PREFIX_TUPLE,
    LINK_STATUS_DOWN,
    LINK_STATUS_UNKNOWN,
    LINK_STATUS_UP,
    VISIBILITY_LIST,
)
from app.utils.datetime_helper import format_local_datetime, to_local_time, utc_now
from app.utils.text_helper import get_link_status_label, get_visibility_label

LOGO_PATH = Path(__file__).resolve().parent.parent / "static" / "images" / "spindo_logo.png"
LOGO_HEIGHT_PX = 38
REPORT_TITLE = "Access Link Register"
REPORT_SUBTITLE = "ALR · ICT"
CONFIDENTIAL_TEXT = "RAHASIA · HANYA UNTUK INTERNAL"
NOTE_TEXT = "Catatan: access note/password tidak ikut diekspor. Waktu dalam WIB."

# kolom sheet Detail
DETAIL_COLUMN_LIST = [
    ("No", 6, True),
    ("Judul", 30, False),
    ("Kategori", 26, False),
    ("URL", 36, False),
    ("Address", 18, False),
    ("Port", 8, True),
    ("Username", 18, False),
    ("Deskripsi", 40, False),
    ("Visibilitas", 12, True),
    ("Status Link", 13, True),
    ("Lampiran", 10, True),
    ("Dibuat Oleh", 22, False),
    ("Dibuat", 18, True),
    ("Diubah", 18, True),
]
DETAIL_URL_COLUMN = 4
DETAIL_STATUS_COLUMN = 10
DETAIL_DATETIME_COLUMN_SET = {13, 14}
DETAIL_HEADER_ROW = 6

# sheet Ringkasan: kolom A cuma margin, isinya di B-F
SUMMARY_COLUMN_WIDTH_LIST = [2, 32, 16, 16, 16, 16]
SUMMARY_META_START_ROW = 5

# warna ngikut brand Spindo
COLOR_DARK = "262626"
COLOR_RED = "C8102E"
COLOR_GREEN = "1E7B34"
COLOR_LINK = "0B5CAD"
COLOR_GRAY_TEXT = "6B6B6B"
COLOR_TILE = "F7F7F7"
COLOR_BAND = "F2F2F2"
COLOR_LINE = "D9D9D9"
STATUS_COLOR_DICT = {LINK_STATUS_UP: COLOR_GREEN, LINK_STATUS_DOWN: COLOR_RED, LINK_STATUS_UNKNOWN: COLOR_GRAY_TEXT}

FONT_NAME = "Calibri"
THIN_LINE = Side(style="thin", color=COLOR_LINE)
CELL_BORDER = Border(left=THIN_LINE, right=THIN_LINE, top=THIN_LINE, bottom=THIN_LINE)
ACCENT_BORDER = Border(bottom=Side(style="medium", color=COLOR_RED))
TOTAL_BORDER = Border(top=Side(style="thin", color=COLOR_DARK), bottom=Side(style="thin", color=COLOR_DARK))
HEADER_FILL = PatternFill("solid", fgColor=COLOR_DARK)
BAND_FILL = PatternFill("solid", fgColor=COLOR_BAND)
TILE_FILL = PatternFill("solid", fgColor=COLOR_TILE)
# [$-421] = locale Indonesia, jadi bulannya tampil "Okt", "Des", dst
DATETIME_FORMAT = "[$-421]dd mmm yyyy hh:mm"
PERCENT_FORMAT = '0"%"'
DETAIL_CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="top", wrap_text=False, indent=0)
DETAIL_TEXT_ALIGNMENT = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)

@lru_cache(maxsize=64)
def make_font(size=10, is_bold=False, color=COLOR_DARK, is_italic=False, is_underline=False):
    """Font seragam satu laporan."""
    return Font(
        name=FONT_NAME, size=size, bold=is_bold, italic=is_italic, color=color,
        underline="single" if is_underline else None,
    )

def is_formula_like(value):
    """Teks yg diawali = + - @ bisa kebaca rumus sama Excel (formula injection)."""
    return isinstance(value, str) and value.startswith(EXPORT_FORMULA_PREFIX_TUPLE)

def write_cell(sheet, row, column, value, font=None, alignment=None, fill=None, border=None, number_format=None):
    """Satu pintu buat nulis sel. Teks mirip rumus dipaksa jadi teks biasa, tampilannya tetep asli tanpa tanda '."""
    cell = sheet.cell(row=row, column=column, value=value)
    if font:
        cell.font = font
    if alignment:
        cell.alignment = alignment
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    if number_format:
        cell.number_format = number_format
    if is_formula_like(value):
        cell.data_type = "s"
        cell.quotePrefix = True
    return cell

def to_excel_time(value):
    """Excel ga ngerti zona waktu, jadi diubah ke WIB terus info zonanya dibuang."""
    local_time = to_local_time(value)
    return local_time.replace(tzinfo=None) if local_time else None

def write_logo(sheet, anchor):
    """Logo Spindo. Kalau file logonya ga ada, dilewatin aja."""
    if not LOGO_PATH.exists():
        return
    logo = Image(str(LOGO_PATH))
    scale = LOGO_HEIGHT_PX / logo.height
    logo.height = LOGO_HEIGHT_PX
    logo.width = int(logo.width * scale)
    sheet.add_image(logo, anchor)

def write_letterhead(sheet, first_column, last_column, title, subtitle):
    """Kop dokumen: logo, judul, sub judul, label rahasia di kanan, garis merah di bawahnya."""
    sheet.row_dimensions[1].height = 32
    write_logo(sheet, f"{get_column_letter(first_column)}1")
    sheet.row_dimensions[2].height = 26
    write_cell(sheet, 2, first_column, title, font=make_font(18, is_bold=True))
    write_cell(sheet, 3, first_column, subtitle, font=make_font(10, color=COLOR_GRAY_TEXT))
    write_cell(
        sheet, 3, last_column, CONFIDENTIAL_TEXT,
        font=make_font(9, is_bold=True, color=COLOR_RED), alignment=Alignment(horizontal="right"),
    )
    for column in range(first_column, last_column + 1):
        sheet.cell(row=3, column=column).border = ACCENT_BORDER

def write_section_title(sheet, row, column, text):
    """Judul bagian kecil di Ringkasan."""
    sheet.row_dimensions[row].height = 22
    write_cell(sheet, row, column, text, font=make_font(12, is_bold=True), alignment=Alignment(vertical="bottom"))

def write_table_header(sheet, row, first_column, title_list):
    """Header tabel gelap, huruf putih."""
    sheet.row_dimensions[row].height = 22
    for offset, title in enumerate(title_list):
        write_cell(
            sheet, row, first_column + offset, title,
            font=make_font(10, is_bold=True, color="FFFFFF"), fill=HEADER_FILL, border=CELL_BORDER,
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        )

def setup_print(sheet, last_column, last_row, orientation, title_row=None):
    """Siap print A4: muat 1 halaman lebar, margin rapi, footer rahasia + nomor halaman."""
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = sheet.page_margins.right = 0.4
    sheet.page_margins.top = sheet.page_margins.bottom = 0.5
    sheet.print_area = f"A1:{get_column_letter(last_column)}{last_row}"
    if title_row:
        sheet.print_title_rows = f"{title_row}:{title_row}"
    sheet.oddFooter.left.text = f"ICT · {CONFIDENTIAL_TEXT.title()}"
    sheet.oddFooter.left.size = 8
    sheet.oddFooter.right.text = "Halaman &P dari &N"
    sheet.oddFooter.right.size = 8

def build_summary_dict(entry_list):
    """Hitung angka buat Ringkasan: per status, per kategori utama (+ status), per visibilitas."""
    status_counter = Counter(entry.status for entry in entry_list)
    visibility_counter = Counter(entry.visibility for entry in entry_list)
    category_status_dict = {}
    for entry in entry_list:
        root_name = get_root_category(entry.category).name if entry.category else "-"
        category_status_dict.setdefault(root_name, Counter())[entry.status] += 1
    category_row_list = sorted(
        category_status_dict.items(), key=lambda item: (-sum(item[1].values()), item[0].lower()),
    )
    return {
        "total": len(entry_list),
        "status_counter": status_counter,
        "visibility_counter": visibility_counter,
        "category_row_list": category_row_list,
    }

def write_meta_block(sheet, start_row, user, filter_text_list, total):
    """Info laporan: siapa yg narik, kapan, filternya apa, berapa data."""
    meta_list = [
        ("Diekspor oleh", f"{user.full_name} ({user.email})"),
        ("Waktu ekspor", to_excel_time(utc_now())),
        ("Filter", ", ".join(filter_text_list) or "Semua data yang dapat diakses"),
        ("Jumlah data", f"{total} link"),
    ]
    for offset, (label, value) in enumerate(meta_list):
        row = start_row + offset
        write_cell(sheet, row, 2, label, font=make_font(10, is_bold=True, color=COLOR_GRAY_TEXT))
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        write_cell(
            sheet, row, 3, value, font=make_font(10), alignment=Alignment(horizontal="left"),
            number_format=DATETIME_FORMAT if label == "Waktu ekspor" else None,
        )
    return start_row + len(meta_list)

def write_kpi_tiles(sheet, row, summary_dict):
    """4 kotak angka besar: total, aktif, tidak aktif, belum dicek (+ persennya)."""    
    total = summary_dict["total"]
    status_counter = summary_dict["status_counter"]
    tile_list = [("Total Link", total, COLOR_DARK, None)]
    for status in (LINK_STATUS_UP, LINK_STATUS_DOWN, LINK_STATUS_UNKNOWN):
        tile_list.append((get_link_status_label(status), status_counter[status], STATUS_COLOR_DICT[status], status))
    sheet.row_dimensions[row].height = 34
    sheet.row_dimensions[row + 1].height = 18
    for offset, (label, value, color, status) in enumerate(tile_list):
        column = 2 + offset
        tile_border = Border(left=THIN_LINE, right=THIN_LINE, top=Side(style="thick", color=color))
        write_cell(
            sheet, row, column, value, font=make_font(22, is_bold=True, color=color), fill=TILE_FILL,
            border=tile_border, alignment=Alignment(horizontal="center", vertical="center"),
        )
        caption = label if status is None else f"{label} · {calculate_percent(value, total)}%"
        write_cell(
            sheet, row + 1, column, caption, font=make_font(9, color=COLOR_GRAY_TEXT), fill=TILE_FILL,
            border=Border(left=THIN_LINE, right=THIN_LINE, bottom=THIN_LINE),
            alignment=Alignment(horizontal="center", vertical="top"),
        )
    return row + 2

def write_category_table(sheet, row, summary_dict):
    """Tabel per kategori utama: jumlah + rincian status, ditutup baris Total."""
    status_order_list = [LINK_STATUS_UP, LINK_STATUS_DOWN, LINK_STATUS_UNKNOWN]
    write_table_header(sheet, row, 2, ["Kategori Utama", "Jumlah"] + [get_link_status_label(status) for status in status_order_list])
    for offset, (category_name, status_counter) in enumerate(summary_dict["category_row_list"]):
        current_row = row + 1 + offset
        fill = BAND_FILL if offset % 2 == 1 else None
        value_list = [category_name, sum(status_counter.values())] + [status_counter[s] for s in status_order_list]
        for column_offset, value in enumerate(value_list):
            write_cell(
                sheet, current_row, 2 + column_offset, value, font=make_font(10), fill=fill, border=CELL_BORDER,
                alignment=Alignment(horizontal="left" if column_offset == 0 else "center", indent=1 if column_offset == 0 else 0),
            )
    total_row = row + 1 + len(summary_dict["category_row_list"])
    total_value_list = ["Total", summary_dict["total"]] + [summary_dict["status_counter"][s] for s in status_order_list]
    for column_offset, value in enumerate(total_value_list):
        write_cell(
            sheet, total_row, 2 + column_offset, value, font=make_font(10, is_bold=True), border=TOTAL_BORDER,
            alignment=Alignment(horizontal="left" if column_offset == 0 else "center", indent=1 if column_offset == 0 else 0),
        )
    return total_row + 1

def write_visibility_table(sheet, row, summary_dict):
    """Tabel per visibilitas: jumlah + persen."""
    write_table_header(sheet, row, 2, ["Visibilitas", "Jumlah", "Persentase"])
    total = summary_dict["total"]
    for offset, visibility in enumerate(VISIBILITY_LIST):
        current_row = row + 1 + offset
        count = summary_dict["visibility_counter"][visibility]
        write_cell(sheet, current_row, 2, get_visibility_label(visibility), font=make_font(10), border=CELL_BORDER,
                   alignment=Alignment(horizontal="left", indent=1))
        write_cell(sheet, current_row, 3, count, font=make_font(10), border=CELL_BORDER,
                   alignment=Alignment(horizontal="center"))
        write_cell(sheet, current_row, 4, calculate_percent(count, total), font=make_font(10), border=CELL_BORDER,
                   alignment=Alignment(horizontal="center"), number_format=PERCENT_FORMAT)
    return row + 1 + len(VISIBILITY_LIST)

def build_summary_sheet(sheet, entry_list, user, filter_text_list):
    """Sheet pertama yg dibuka: ringkasan buat dibaca sekilas (pimpinan), detailnya di sheet sebelah."""
    sheet.title = "Ringkasan"
    sheet.sheet_view.showGridLines = False
    for column, width in enumerate(SUMMARY_COLUMN_WIDTH_LIST, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    last_column = len(SUMMARY_COLUMN_WIDTH_LIST)
    summary_dict = build_summary_dict(entry_list)

    write_letterhead(sheet, 2, last_column, REPORT_TITLE, REPORT_SUBTITLE)
    row = write_meta_block(sheet, SUMMARY_META_START_ROW, user, filter_text_list, summary_dict["total"]) + 1
    write_section_title(sheet, row, 2, "Status Link")
    row = write_kpi_tiles(sheet, row + 1, summary_dict) + 1
    write_section_title(sheet, row, 2, "Per Kategori Utama")
    row = write_category_table(sheet, row + 1, summary_dict) + 1
    write_section_title(sheet, row, 2, "Per Visibilitas")
    row = write_visibility_table(sheet, row + 1, summary_dict) + 1
    write_cell(sheet, row, 2, "Rincian lengkap tiap link ada di sheet \"Detail\".", font=make_font(9, is_italic=True, color=COLOR_GRAY_TEXT))
    write_cell(sheet, row + 1, 2, NOTE_TEXT, font=make_font(9, is_italic=True, color=COLOR_GRAY_TEXT))
    setup_print(sheet, last_column, row + 1, "portrait")

def build_detail_row(row_number, entry):
    """Satu baris detail buat satu link. Access note sengaja ga ada."""
    return [
        row_number,
        entry.title,
        build_category_label(entry.category),
        entry.url,
        entry.address,
        entry.port,
        entry.username,
        entry.description,
        get_visibility_label(entry.visibility),
        get_link_status_label(entry.status),
        entry.attachment_count,
        entry.owner.full_name,
        to_excel_time(entry.created_at),
        to_excel_time(entry.updated_at),
    ]

def write_detail_row(sheet, row, value_list, entry, is_band_row):
    """Tulis satu baris detail: garis tipis, selang-seling abu, URL bisa diklik, status berwarna."""
    fill = BAND_FILL if is_band_row else None
    for column, value in enumerate(value_list, start=1):
        is_center = DETAIL_COLUMN_LIST[column - 1][2]
        font = make_font(10)
        if column == DETAIL_STATUS_COLUMN:
            font = make_font(10, is_bold=True, color=STATUS_COLOR_DICT.get(entry.status, COLOR_DARK))
        cell = write_cell(
            sheet, row, column, value if value not in ("", None) else None, font=font, fill=fill, border=CELL_BORDER,
            alignment=DETAIL_CENTER_ALIGNMENT if is_center else DETAIL_TEXT_ALIGNMENT,
            number_format=DATETIME_FORMAT if column in DETAIL_DATETIME_COLUMN_SET else None,
        )
        if column == DETAIL_URL_COLUMN and isinstance(value, str) and value.startswith(("http://", "https://")):
            cell.hyperlink = value
            cell.font = make_font(10, color=COLOR_LINK, is_underline=True)

def build_detail_sheet(sheet, entry_list, user, filter_text_list):
    """Sheet kedua: tabel lengkap tiap link, bisa difilter & di-print."""
    sheet.title = "Detail"
    sheet.sheet_view.showGridLines = False
    for column, (_, width, _) in enumerate(DETAIL_COLUMN_LIST, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    last_column = len(DETAIL_COLUMN_LIST)
    last_column_letter = get_column_letter(last_column)

    subtitle = f"{REPORT_SUBTITLE}  |  Diekspor {format_local_datetime(utc_now())} oleh {user.full_name}"
    write_letterhead(sheet, 1, last_column, "Detail Access Link Register", subtitle)
    write_cell(
        sheet, 4, 1, f"Filter: {', '.join(filter_text_list) or 'Semua data yang dapat diakses'}",
        font=make_font(9, color=COLOR_GRAY_TEXT),
    )
    write_table_header(sheet, DETAIL_HEADER_ROW, 1, [title for title, _, _ in DETAIL_COLUMN_LIST])

    for offset, entry in enumerate(entry_list):
        value_list = build_detail_row(offset + 1, entry)
        write_detail_row(sheet, DETAIL_HEADER_ROW + 1 + offset, value_list, entry, offset % 2 == 1)

    last_row = DETAIL_HEADER_ROW + len(entry_list)
    if not entry_list:
        sheet.merge_cells(start_row=last_row + 1, start_column=1, end_row=last_row + 1, end_column=last_column)
        write_cell(sheet, last_row + 1, 1, "Tidak ada data sesuai filter.", font=make_font(10, is_italic=True, color=COLOR_GRAY_TEXT),
                   alignment=Alignment(horizontal="center"))
        last_row += 1
    sheet.auto_filter.ref = f"A{DETAIL_HEADER_ROW}:{last_column_letter}{max(last_row, DETAIL_HEADER_ROW)}"
    sheet.freeze_panes = f"C{DETAIL_HEADER_ROW + 1}"
    write_cell(sheet, last_row + 2, 1, NOTE_TEXT, font=make_font(9, is_italic=True, color=COLOR_GRAY_TEXT))
    setup_print(sheet, last_column, last_row + 2, "landscape", title_row=DETAIL_HEADER_ROW)

def build_entry_workbook(entry_list, user, filter_text_list):
    """Laporan 2 sheet: Ringkasan (dibuka pertama) + Detail. Return BytesIO siap dikirim."""
    workbook = Workbook()
    workbook.properties.title = REPORT_TITLE
    workbook.properties.subject = REPORT_SUBTITLE
    workbook.properties.creator = f"ICT · {user.full_name}"
    build_summary_sheet(workbook.active, entry_list, user, filter_text_list)
    build_detail_sheet(workbook.create_sheet(), entry_list, user, filter_text_list)
    workbook.active = 0

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
