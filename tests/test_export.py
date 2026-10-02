"""Test export Excel daftar link (SR-10): hak akses, filter, access note, formula injection, audit."""
from io import BytesIO
from openpyxl import load_workbook
from app.extensions import db
from app.models import AuditLog
from app.services.access_entry_service import create_access_entry, list_visible_entries_for_export
from app.utils.constants import AUDIT_ACTION_EXPORT, VISIBILITY_PRIVATE, VISIBILITY_PUBLIC, XLSX_MIMETYPE

def build_entry_dict(category, title, url, **override_dict):
    """Helper: data link contoh."""
    entry_dict = {
        "category_id": category.id, "title": title, "url": url,
        "address": "", "port": "", "username": "", "access_note": "", "description": "",
        "visibility": VISIBILITY_PRIVATE,
    }
    entry_dict.update(override_dict)
    return entry_dict

def load_report(response):
    """Buka file Excel dari response."""
    return load_workbook(BytesIO(response.data))

def read_sheet_value_list(response, sheet_name="Detail"):
    """Semua isi sel satu sheet jadi list baris."""
    return [list(row) for row in load_report(response)[sheet_name].iter_rows(values_only=True)]

def read_report_text(response):
    """Isi semua sheet digabung jadi teks (buat ngecek ada/ngga-nya sesuatu)."""
    workbook = load_report(response)
    return str([list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets])

def read_table_row_list(response):
    """Baris tabelnya aja: mulai dari judul kolom "No" sampe baris data terakhir (kop & catatan dibuang)."""
    value_list = read_sheet_value_list(response)
    header_index = next(index for index, row in enumerate(value_list) if row[0] == "No")
    table_row_list = [value_list[header_index]]
    for row in value_list[header_index + 1:]:
        if not isinstance(row[0], int):
            break
        table_row_list.append(row)
    return table_row_list

def read_title_list(response):
    """Kolom Judul (kolom ke-2) tanpa header."""
    return [row[1] for row in read_table_row_list(response)[1:]]

def test_export_requires_login(client):
    """Negative: belum login ga bisa export."""
    response = client.get("/export")
    assert response.status_code == 302
    assert "/auth/login" in response.headers["Location"]

def test_export_returns_xlsx_file(logged_in_client, registered_user, category_dict):
    """Positive: file kekirim sebagai xlsx dengan nama Laporan_Access_Link_*.xlsx + 2 sheet."""
    create_access_entry(registered_user, build_entry_dict(category_dict["Web"], "Portal HR", "https://hr.spindo.com"))
    response = logged_in_client.get("/export")
    assert response.status_code == 200
    assert response.mimetype == XLSX_MIMETYPE
    disposition_text = response.headers["Content-Disposition"]
    assert "attachment" in disposition_text
    assert "Laporan_Access_Link_" in disposition_text and ".xlsx" in disposition_text
    assert load_report(response).sheetnames == ["Ringkasan", "Detail"]
    assert read_table_row_list(response)[0][:2] == ["No", "Judul"]

def test_export_only_visible_entries(logged_in_client, registered_user, other_user, category_dict):
    """Positive/Negative: private orang lain ga ikut, private sendiri & public orang lain ikut."""
    web = category_dict["Web"]
    create_access_entry(registered_user, build_entry_dict(web, "Private Saya", "https://saya.spindo.com"))
    create_access_entry(other_user, build_entry_dict(web, "Private Orang", "https://orang.spindo.com"))
    create_access_entry(other_user, build_entry_dict(
        web, "Public Orang", "https://public.spindo.com", visibility=VISIBILITY_PUBLIC,
    ))
    title_list = read_title_list(logged_in_client.get("/export"))
    assert "Private Saya" in title_list
    assert "Public Orang" in title_list
    assert "Private Orang" not in title_list

def test_export_follows_filter(logged_in_client, registered_user, category_dict):
    """Positive: filter yg sama kayak di Home ikut kepake."""
    create_access_entry(registered_user, build_entry_dict(category_dict["Web"], "Portal HR", "https://hr.spindo.com"))
    create_access_entry(registered_user, build_entry_dict(
        category_dict["Network"], "Switch Core", "", address="10.0.0.2", port="22",
    ))
    response = logged_in_client.get(f"/export?category_id={category_dict['Network'].id}&q=switch")
    assert read_title_list(response) == ["Switch Core"]
    report_text = read_report_text(response)
    assert 'Kata kunci "switch"' in report_text
    assert "Kategori Network" in report_text

def test_export_never_contains_access_note(logged_in_client, registered_user, category_dict):
    """Negative: access note (password) ga boleh nongol di sel manapun."""
    secret_text = "RahasiaBanget123"
    create_access_entry(registered_user, build_entry_dict(
        category_dict["Web"], "Portal HR", "https://hr.spindo.com", access_note=secret_text,
    ))
    response = logged_in_client.get("/export")
    assert secret_text not in read_report_text(response)

def test_export_escapes_formula(logged_in_client, registered_user, category_dict):
    """Negative: teks yg mirip rumus Excel disimpen sebagai teks biasa, bukan rumus."""
    create_access_entry(registered_user, build_entry_dict(
        category_dict["Web"], "=1+1", "https://hr.spindo.com", description="+cmd",
    ))
    sheet = load_report(logged_in_client.get("/export"))["Detail"]
    header_row = next(row[0].row for row in sheet.iter_rows() if row[0].value == "No")
    title_cell = sheet.cell(row=header_row + 1, column=2)
    description_cell = sheet.cell(row=header_row + 1, column=8)
    assert title_cell.value == "=1+1"
    assert title_cell.data_type == "s"
    assert description_cell.data_type == "s"

def test_export_summary_counts(logged_in_client, registered_user, other_user, category_dict):
    """Positive: angka di Ringkasan sesuai data yg boleh diliat (private orang lain ga kehitung)."""
    web = category_dict["Web"]
    create_access_entry(registered_user, build_entry_dict(web, "Portal HR", "https://hr.spindo.com"))
    create_access_entry(registered_user, build_entry_dict(
        category_dict["Network"], "Switch Core", "", address="10.0.0.2", port="22", visibility=VISIBILITY_PUBLIC,
    ))
    create_access_entry(other_user, build_entry_dict(web, "Private Orang", "https://orang.spindo.com"))
    value_list = read_sheet_value_list(logged_in_client.get("/export"), "Ringkasan")
    category_row_dict = {row[1]: row[2] for row in value_list if row[1] in ("Web", "Network", "Total")}
    assert category_row_dict == {"Web": 1, "Network": 1, "Total": 2}

def test_export_is_audited(logged_in_client, registered_user, category_dict):
    """Positive: tiap export kecatat di audit log + jumlah datanya."""
    create_access_entry(registered_user, build_entry_dict(category_dict["Web"], "Portal HR", "https://hr.spindo.com"))
    logged_in_client.get("/export")
    audit_log = db.session.execute(
        db.select(AuditLog).where(AuditLog.action == AUDIT_ACTION_EXPORT)
    ).scalar_one()
    assert audit_log.user_id == registered_user.id
    assert audit_log.new_data["row_count"] == 1

def test_export_list_is_truncated(registered_user, category_dict):
    """Edge: data lebih dari batas dipotong + ditandain kepotong."""
    web = category_dict["Web"]
    create_access_entry(registered_user, build_entry_dict(web, "Link Satu", "https://satu.spindo.com"))
    create_access_entry(registered_user, build_entry_dict(web, "Link Dua", "https://dua.spindo.com"))
    entry_list, is_truncated = list_visible_entries_for_export(registered_user, max_count=1)
    assert len(entry_list) == 1
    assert is_truncated is True
    entry_list, is_truncated = list_visible_entries_for_export(registered_user, max_count=5)
    assert len(entry_list) == 2
    assert is_truncated is False

def test_home_shows_export_button(logged_in_client, registered_user, category_dict):
    """Positive: tombol Export Excel muncul kalau ada data, link-nya bawa filter."""
    create_access_entry(registered_user, build_entry_dict(category_dict["Web"], "Portal HR", "https://hr.spindo.com"))
    html_text = logged_in_client.get("/?q=portal").get_data(as_text=True)
    assert "Export Excel" in html_text
    assert "/export?q=portal" in html_text

def test_home_hides_export_button_when_empty(logged_in_client, category_dict):
    """Negative: ga ada data, tombolnya ga usah muncul."""
    html_text = logged_in_client.get("/").get_data(as_text=True)
    assert "Export Excel" not in html_text