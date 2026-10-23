"""Test import link dari Excel: semua-atau-tidak-sama-sekali, aturan sama kayak form Add Link, file berbahaya ditolak."""
import io

from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models import AccessEntry, AuditLog, Category
from app.security.encryption_service import decrypt_credential
from app.services import import_service
from app.services.access_entry_service import create_access_entry
from app.utils.constants import AUDIT_ACTION_CREATE, AUDIT_ACTION_IMPORT
from tests.conftest import build_office_zip_bytes
from tests.test_entry_routes import build_service_dict

HEADER_LIST = ["Title", "Category", "URL", "Address", "Port", "Username", "Access Note", "Description", "Visibility"]


def build_xlsx(row_list, header_list=HEADER_LIST):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header_list)
    for row in row_list:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def post_import(client, content_bytes, filename="links.xlsx"):
    return client.post("/entries/import", data={"import_file": (io.BytesIO(content_bytes), filename)},
                       content_type="multipart/form-data")


def count_entry():
    return db.session.execute(db.select(db.func.count(AccessEntry.id))).scalar()


def web_row(title, url, **override):
    row = {"Title": title, "Category": "Web", "URL": url, "Address": "", "Port": "", "Username": "",
           "Access Note": "", "Description": "", "Visibility": ""}
    row.update(override)
    return [row[header] for header in HEADER_LIST]


def test_template_download_has_columns_and_categories(logged_in_client, category_dict):
    """Positive: template berisi judul kolom yg dipake import + daftar kategori aktif."""
    sap = Category(parent_id=category_dict["Web"].id, name="SAP", is_active=True)
    db.session.add(sap)
    db.session.commit()
    response = logged_in_client.get("/entries/import/template")
    assert response.status_code == 200 and "attachment" in response.headers["Content-Disposition"]
    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.sheetnames == ["Links", "Guide", "Categories"]
    header_list = [cell.value for cell in workbook["Links"][5]]
    assert header_list == ["Title *", "Category *"] + HEADER_LIST[2:]
    category_list = [row[0] for row in workbook["Categories"].iter_rows(min_row=6, values_only=True)]
    assert "Web" in category_list and "Web › SAP" in category_list
    validation_range_list = [str(validation.sqref) for validation in workbook["Links"].data_validations.dataValidation]
    assert "B6:B505" in validation_range_list and "I6:I505" in validation_range_list


def test_filled_official_template_imports(logged_in_client, registered_user, category_dict):
    """Positive: template resmi (header di bawah kop) diisi pake nilai dropdown -> ke-import.
    Visibility Public dari user biasa disimpen Private (Public cuma admin)."""
    sap = Category(parent_id=category_dict["Web"].id, name="SAP", is_active=True)
    db.session.add(sap)
    db.session.commit()
    workbook = load_workbook(io.BytesIO(logged_in_client.get("/entries/import/template").data))
    sheet = workbook["Links"]
    for column_index, value in enumerate(["SAP Prod", "Web › SAP", "https://sap.spindo.com", "", None, "", "", "", "Public"], start=1):
        sheet.cell(row=6, column=column_index, value=value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = post_import(logged_in_client, buffer.getvalue(), "ALR_import_template.xlsx")
    assert response.status_code == 302, response.get_data(as_text=True)[:2000]
    entry = db.session.execute(db.select(AccessEntry)).scalar_one()
    assert (entry.title, entry.category_id, entry.visibility, entry.user_id) == ("SAP Prod", sap.id, "private", registered_user.id)


def test_empty_official_template_has_nothing_to_import(logged_in_client, category_dict):
    """Negative: template belum diisi (ga ada baris contoh yg ikut ke-import) -> pesan jelas, ga ada yg kesimpen."""
    template_bytes = logged_in_client.get("/entries/import/template").data
    assert "no links to import" in post_import(logged_in_client, template_bytes).get_data(as_text=True)
    assert count_entry() == 0


def test_import_creates_all_rows_with_encrypted_note(logged_in_client, registered_user, category_dict):
    """Positive: semua baris bener -> kesimpen semua (milik yg upload), password dienkripsi, kecatat audit."""
    sap = Category(parent_id=category_dict["Web"].id, name="SAP", is_active=True)
    db.session.add(sap)
    db.session.commit()
    content = build_xlsx([
        web_row("Portal HR", "https://hr.spindo.com", **{"Access Note": "Rahasia123", "Visibility": "Public"}),
        web_row("SAP Prod", "https://sap.spindo.com", Category="web > sap"),
        ["Switch Core", "Network", "", "10.0.0.2", 22, "admin", "", "", "private"],
    ])
    response = post_import(logged_in_client, content)
    assert response.status_code == 302 and "/entries/" in response.location
    entry_list = db.session.execute(db.select(AccessEntry).order_by(AccessEntry.id)).scalars().all()
    assert [(entry.title, entry.user_id) for entry in entry_list] == [
        ("Portal HR", registered_user.id), ("SAP Prod", registered_user.id), ("Switch Core", registered_user.id),
    ]
    # baris Public dari user biasa ikut Private (Public cuma admin)
    assert entry_list[0].visibility == "private" and entry_list[2].visibility == "private"
    assert entry_list[1].category_id == sap.id and entry_list[2].port == 22
    assert entry_list[0].encrypted_access_note != "Rahasia123"
    assert decrypt_credential(entry_list[0].encrypted_access_note) == "Rahasia123"
    import_log = db.session.execute(db.select(AuditLog).filter_by(action=AUDIT_ACTION_IMPORT)).scalar_one()
    assert import_log.new_data == {"row_count": 3, "filename": "links.xlsx"}
    assert db.session.execute(db.select(db.func.count(AuditLog.id)).filter_by(action=AUDIT_ACTION_CREATE)).scalar() >= 3


def test_one_bad_row_imports_nothing(logged_in_client, category_dict):
    """Negative: 1 baris salah (Web tanpa URL) -> ga ada yg disimpen, error nyebut nomor baris Excel."""
    content = build_xlsx([web_row("Bagus", "https://ok.spindo.com"), web_row("Rusak", "")])
    response = post_import(logged_in_client, content)
    html_text = response.get_data(as_text=True)
    assert response.status_code == 200 and "Nothing was imported" in html_text
    assert "Row 3: URL is required" in html_text
    assert count_entry() == 0


def test_duplicate_inside_file_and_against_database(logged_in_client, registered_user, category_dict):
    """Negative: URL kembar antar baris & URL yg udah ada di database ditolak, per baris."""
    create_access_entry(registered_user, build_service_dict(category_dict["Web"], url="https://lama.spindo.com"))
    content = build_xlsx([
        web_row("A", "https://baru.spindo.com"),
        web_row("B", "https://baru.spindo.com"),
        web_row("C", "https://lama.spindo.com"),
    ])
    html_text = post_import(logged_in_client, content).get_data(as_text=True)
    assert "Row 3: Same URL as row 2" in html_text
    assert "Row 4: Already registered in link" in html_text
    assert count_entry() == 1


def test_unknown_category_and_bad_visibility(logged_in_client, category_dict):
    """Negative: kategori ga dikenal / visibilitas aneh -> error jelas."""
    content = build_xlsx([web_row("A", "https://a.spindo.com", Category="Ngasal"),
                          web_row("B", "https://b.spindo.com", Visibility="Rahasia")])
    html_text = post_import(logged_in_client, content).get_data(as_text=True)
    assert "Row 2: Unknown or inactive category" in html_text
    assert "Row 3: Visibility must be Public or Private" in html_text


def test_inactive_category_rejected(logged_in_client, category_dict):
    """Negative: kategori nonaktif ga bisa dipake import (sama kayak form)."""
    printer = Category(name="Printer", is_active=False)
    db.session.add(printer)
    db.session.commit()
    content = build_xlsx([["Printer Lt 2", "Printer", "", "", "", "", "", "", ""]])
    assert "Unknown or inactive category" in post_import(logged_in_client, content).get_data(as_text=True)


def test_dangerous_or_wrong_files_rejected(logged_in_client, category_dict, sample_file_dict):
    """Negative (security): bukan .xlsx, exe diganti nama, & xlsx bermakro ditolak sebelum dibaca."""
    assert "Upload the .xlsx file" in post_import(logged_in_client, b"a,b", "links.csv").get_data(as_text=True)
    assert "File rejected" in post_import(logged_in_client, sample_file_dict["exe"]).get_data(as_text=True)
    macro_xlsx = build_office_zip_bytes("xl/", has_macro=True)
    assert "File rejected" in post_import(logged_in_client, macro_xlsx).get_data(as_text=True)
    assert count_entry() == 0


def test_missing_columns_and_empty_file(logged_in_client, category_dict):
    """Negative: kolom wajib ga ada / file kosong -> pesan jelas."""
    no_category = build_xlsx([["A", "https://a.spindo.com"]], header_list=["Title", "URL"])
    assert "Missing column(s): Category" in post_import(logged_in_client, no_category).get_data(as_text=True)
    assert "no links to import" in post_import(logged_in_client, build_xlsx([])).get_data(as_text=True)


def test_row_limit(logged_in_client, category_dict, monkeypatch):
    """Negative: lewat batas baris -> ditolak semua (biar server ga berat)."""
    monkeypatch.setattr(import_service, "IMPORT_MAX_ROW_COUNT", 2)
    content = build_xlsx([web_row(f"L{index}", f"https://l{index}.spindo.com") for index in range(3)])
    assert "Up to 2 links per import" in post_import(logged_in_client, content).get_data(as_text=True)
    assert count_entry() == 0


def test_import_requires_login(client):
    """Negative: belum login ga bisa import / download template."""
    assert "/auth/login" in client.get("/entries/import").location
    assert "/auth/login" in client.get("/entries/import/template").location


def test_filled_column_not_used_by_category_is_rejected(logged_in_client, category_dict):
    """Negative (data hilang): password di baris General & URL di baris Network ditolak, bukan dibuang diem-diem."""
    general_row = web_row("Catatan", "", Category="General", **{"Access Note": "rahasia123"})
    network_row = web_row("Switch", "https://switch.spindo.com", Category="Network", Address="10.0.0.9")
    response = post_import(logged_in_client, build_xlsx([general_row, network_row]))
    html_text = response.get_data(as_text=True)
    assert "Row 2: Access Note is not used by category General" in html_text
    assert "Row 3: URL is not used by category Network" in html_text
    assert count_entry() == 0


def test_import_query_count_does_not_grow_per_row(app, registered_user, category_dict):
    """Positive (efisiensi, temuan audit): cek duplikat & simpan dikerjain sekaligus, bukan beberapa query per baris."""
    from sqlalchemy import event
    from werkzeug.datastructures import FileStorage

    def count_statement(row_count):
        content = build_xlsx([web_row(f"Link {row_count}-{index}", f"https://q{row_count}-{index}.spindo.com")
                              for index in range(row_count)])
        counter = [0]
        def on_execute(_conn, _cursor, statement, *_):
            # cuma SELECT yg dihitung: INSERT digabung atau nggak tergantung database (PostgreSQL digabung, SQLite nggak)
            counter[0] += statement.lstrip().upper().startswith("SELECT")
        event.listen(db.engine, "before_cursor_execute", on_execute)
        try:
            import_service.import_entries(registered_user, FileStorage(stream=io.BytesIO(content), filename="x.xlsx"))
        finally:
            event.remove(db.engine, "before_cursor_execute", on_execute)
        return counter[0]

    with app.test_request_context("/"):
        # pemanasan: import pertama ngisi cache kategori, abis commit data user dibaca ulang sekali
        count_statement(1)
        small, large = count_statement(10), count_statement(60)
    # 50 baris tambahan ga boleh nambah query baca sama sekali
    assert large == small, (small, large)


def test_admin_import_keeps_public(admin_client, admin_user, category_dict):
    """Positive: admin boleh import link Public."""
    response = post_import(admin_client, build_xlsx([web_row("Portal Umum", "https://umum.spindo.com", Visibility="Public")]))
    assert response.status_code == 302
    assert db.session.execute(db.select(AccessEntry.visibility)).scalar_one() == "public"
